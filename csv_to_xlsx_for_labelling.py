"""
Run this ONCE, before labelling, to turn data/golden_set.csv into a properly
formatted data/golden_set.xlsx:
  - correct UTF-8 encoding (fixes garbled accented/non-English characters)
  - wide, wrap-text columns so you can actually read customer_text
  - click-to-pick dropdowns for gold_intent and gold_escalate (no more typing/typos)

Run: python csv_to_xlsx_for_labelling.py
Then open data/golden_set.xlsx (NOT the .csv) in Excel/WPS and label there.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment
import config as C

SRC = C.GOLDEN_CSV
OUT = C.GOLDEN_CSV.replace(".csv", ".xlsx")


def main():
    df = pd.read_csv(SRC, encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "golden_set"

    ws.append(list(df.columns))
    for _, row in df.iterrows():
        ws.append(list(row))

    # Column widths tuned for readability: id, customer_text, brand_reply, intent, escalate, notes
    widths = {"A": 16, "B": 70, "C": 60, "D": 30, "E": 14, "F": 40}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=3):
        for cell in row:
            cell.alignment = wrap

    # Freeze header row so it stays visible while scrolling
    ws.freeze_panes = "A2"

    # Dropdown validation: gold_intent (col D) and gold_escalate (col E)
    intent_list = ",".join(C.INTENTS)
    dv_intent = DataValidation(type="list", formula1=f'"{intent_list}"', allow_blank=True)
    ws.add_data_validation(dv_intent)
    dv_intent.add(f"D2:D{ws.max_row}")

    dv_escalate = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    ws.add_data_validation(dv_escalate)
    dv_escalate.add(f"E2:E{ws.max_row}")

    wb.save(OUT)
    print(f"Saved -> {OUT}")
    print("Open THIS file (not the .csv) in Excel/WPS. Columns D and E are dropdowns.")
    print("When done labelling, run: python xlsx_to_csv_after_labelling.py")


if __name__ == "__main__":
    main()
